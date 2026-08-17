import json
from pathlib import Path

from openpyxl import load_workbook

from improve_yourself.excel_report import build_excel_report


def test_portable_excel_report_excludes_non_match_rounds(tmp_path: Path) -> None:
    analysis = tmp_path / "analysis.json"
    analysis.write_text(json.dumps({
        "schema": "iy.analysis/v1", "map_name": "de_anubis", "source_sha256": "a" * 64,
        "kills": [
            {"round_number": 0, "tick": 5, "attacker": "knife", "victim": "other"},
            {"round_number": 1, "tick": 10, "attacker": "player", "victim": "other", "headshot": True, "weapon": "ak47"},
        ],
        "multikills": [{"round_number": 0, "player": "knife", "first_tick": 5}, {"round_number": 1, "player": "player", "first_tick": 10, "last_tick": 12, "kill_count": 2, "victims": ["other"]}],
        "review_hints": [], "available_channels": ["kills"], "data_quality": {"status": "limited"}, "user_view": {"assessment": {"status": "Hinweis"}},
        "review_profile": {"id": "improve-default", "version": "v1", "label": "Improve Default V1", "disclaimer": "Kontext prüfen."},
    }), encoding="utf-8")
    output = build_excel_report(analysis, tmp_path / "report.xlsx")
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["Übersicht", "Runden", "Spieler", "Kills", "Multi-Kills & Szenen", "Review-Hinweise", "Datenqualität", "Info"]
    assert workbook["Runden"]["A3"].value == 1
    assert workbook["Kills"]["A3"].value == 1
    assert workbook["Info"]["B5"].value == "iy.analyzer_excel_report/v1"
    assert workbook["Kills"].freeze_panes == "A3"
    assert "KillsTable" in workbook["Kills"].tables
