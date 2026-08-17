"""Portable, local Analyzer Excel report; no Office, Node or COM runtime."""

from __future__ import annotations

import argparse
import json
from copy import copy
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

EXCEL_REPORT_SCHEMA = "iy.analyzer_excel_report/v1"
_NAVY, _BLUE, _SKY, _PALE, _AMBER, _WHITE, _TEXT = "102A43", "1F5A94", "DCEEFF", "F5F9FD", "FFF2CC", "FFFFFF", "172B4D"
_BORDER = Border(bottom=Side(style="thin", color="C9D7E6"))


def _load(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict) or value.get("schema") != "iy.analysis/v1":
        raise ValueError("expected iy.analysis/v1")
    return value


def _title(sheet: Any, text: str, end: str) -> None:
    sheet.merge_cells(f"A1:{end}1")
    cell = sheet["A1"]; cell.value = text
    cell.fill = PatternFill("solid", fgColor=_NAVY); cell.font = Font(bold=True, color=_WHITE, size=18)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30


def _header(sheet: Any, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=_SKY); cell.font = Font(bold=True, color=_TEXT)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = _BORDER


def _table(sheet: Any, name: str, start: int, end: int, columns: int) -> None:
    if end < start + 1:
        return
    ref = f"A{start}:{chr(64 + columns)}{end}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table); sheet.freeze_panes = f"A{start + 1}"; sheet.auto_filter.ref = ref


def _widths(sheet: Any, values: list[float]) -> None:
    for index, width in enumerate(values, 1):
        sheet.column_dimensions[chr(64 + index)].width = width


def build_excel_report(analysis_path: Path, output_path: Path) -> Path:
    """Export accepted Analyzer facts while excluding non-regular round events."""
    analysis = _load(analysis_path)
    kills = sorted((item for item in analysis.get("kills", []) if isinstance(item, dict) and isinstance(item.get("round_number"), int) and item["round_number"] >= 1), key=lambda item: (item["round_number"], int(item.get("tick", 0))))
    scenes = sorted((item for item in analysis.get("multikills", []) if isinstance(item, dict) and isinstance(item.get("round_number"), int) and item["round_number"] >= 1), key=lambda item: (item["round_number"], int(item.get("first_tick", 0))))
    profile = analysis.get("review_profile") if isinstance(analysis.get("review_profile"), dict) else {}
    hints = [item for item in analysis.get("review_hints", []) if isinstance(item, dict)]
    hints_by_scene = {(item.get("observed_facts", {}).get("round_number"), item.get("observed_facts", {}).get("player"), item.get("observed_facts", {}).get("first_tick")): item for item in hints if isinstance(item.get("observed_facts"), dict)}
    players = sorted({str(name) for kill in kills for name in (kill.get("attacker"), kill.get("victim")) if name})
    rounds = sorted({int(kill["round_number"]) for kill in kills})
    workbook = Workbook(); overview = workbook.active; overview.title = "Übersicht"
    round_sheet = workbook.create_sheet("Runden"); player_sheet = workbook.create_sheet("Spieler"); kill_sheet = workbook.create_sheet("Kills")
    scene_sheet = workbook.create_sheet("Multi-Kills & Szenen"); hint_sheet = workbook.create_sheet("Review-Hinweise"); quality_sheet = workbook.create_sheet("Datenqualität"); info_sheet = workbook.create_sheet("Info")

    _title(overview, "Improve Yourself · Analyzer Report", "H")
    overview.merge_cells("A2:H2"); overview["A2"] = "Lokaler Match- und Szenenbericht. Automatische Hinweise sind keine Cheat-Nachweise."
    overview["A2"].fill = PatternFill("solid", fgColor=_PALE); overview["A2"].alignment = Alignment(wrap_text=True)
    _header(overview, 4, ["Match-Überblick", "Wert"])
    headshots = sum(bool(item.get("headshot")) for item in kills)
    for row, values in enumerate((("Map", analysis.get("map_name") or "Unbekannte Map"), ("Analysierte Runden", len(rounds)), ("Erkannte Kills", len(kills)), ("Review-Szenen", len(scenes)), ("Beteiligte Spieler", len(players)), ("Headshots", headshots)), 5):
        overview.cell(row, 1, values[0]).font = Font(bold=True, color=_TEXT); overview.cell(row, 2, values[1]); overview.cell(row, 1).border = overview.cell(row, 2).border = _BORDER
    overview.merge_cells("A13:H15"); overview["A13"] = "Review-Hinweise und Szenen zeigen nur vorhandene Analyzer-Fakten. Sie sind ein Einstieg für menschliche Prüfung, keine Aussage über Absicht oder Cheating."
    overview["A13"].fill = PatternFill("solid", fgColor=_AMBER); overview["A13"].alignment = Alignment(wrap_text=True, vertical="top")
    _widths(overview, [25, 18, 14, 14, 14, 22, 22, 22])

    _title(round_sheet, "Runden · vorhandene Ereignisübersicht", "C"); _header(round_sheet, 2, ["Runde", "Erkannte Kills", "Multi-Kill-Szenen"])
    kill_counts, scene_counts = Counter(item["round_number"] for item in kills), Counter(item["round_number"] for item in scenes)
    for row, number in enumerate(rounds, 3): round_sheet.append([number, kill_counts[number], scene_counts[number]])
    _table(round_sheet, "RoundsTable", 2, len(rounds) + 2, 3); _widths(round_sheet, [18, 20, 22])

    _title(player_sheet, "Spieler · verlässliche Kill-Ereignis-Übersicht", "F"); _header(player_sheet, 2, ["Spieler", "Kills", "Tode", "Headshots", "Headshot-Rate", "Multi-Kill-Szenen"])
    for name in players:
        player_kills = [item for item in kills if item.get("attacker") == name]; deaths = sum(item.get("victim") == name for item in kills); hs = sum(bool(item.get("headshot")) for item in player_kills)
        player_sheet.append([name, len(player_kills), deaths, hs, hs / len(player_kills) if player_kills else 0, sum(item.get("player") == name for item in scenes)])
    for row in range(3, len(players) + 3): player_sheet.cell(row, 5).number_format = "0.0%"
    _table(player_sheet, "PlayersTable", 2, len(players) + 2, 6); _widths(player_sheet, [24, 14, 14, 15, 18, 20])

    _title(kill_sheet, "Kills · vorhandene Ereignisdaten", "F"); _header(kill_sheet, 2, ["Runde", "Tick", "Killer", "Opfer", "Waffe", "Headshot"])
    for item in kills: kill_sheet.append([item["round_number"], item.get("tick"), item.get("attacker"), item.get("victim"), item.get("weapon") or "Unbekannt", "Ja" if item.get("headshot") else "Nein"])
    _table(kill_sheet, "KillsTable", 2, len(kills) + 2, 6); _widths(kill_sheet, [14, 16, 24, 24, 18, 14])

    _title(scene_sheet, "Multi-Kills & Szenen · Review-Einstieg", "J"); _header(scene_sheet, 2, ["Runde", "Spieler", "Kills", "Erster Tick", "Letzter Tick", "Beteiligte Opfer", "Bewertungskategorie", "Auslösendes Kriterium", "Schwelle", "Beobachtete Fakten"])
    for scene in scenes:
        hint = hints_by_scene.get((scene.get("round_number"), scene.get("player"), scene.get("first_tick"))) or {}; threshold = hint.get("threshold") if isinstance(hint.get("threshold"), dict) else {}
        scene_sheet.append([scene.get("round_number"), scene.get("player"), scene.get("kill_count"), scene.get("first_tick"), scene.get("last_tick"), ", ".join(map(str, scene.get("victims") or [])), hint.get("category") or "vorhandener Marker", hint.get("criterion_label") or "Kein Profilhinweis", f"{threshold.get('field', '—')} {threshold.get('operator', '')} {threshold.get('value', '')}".strip(), hint.get("message") or "Vorhandener Multi-Kill-Marker; im Spielkontext prüfen."])
    _table(scene_sheet, "ScenesTable", 2, len(scenes) + 2, 10); _widths(scene_sheet, [12, 22, 12, 16, 16, 42, 24, 28, 22, 48])

    _title(hint_sheet, "Review-Hinweise · Spieler → Runde → Tick", "H"); _header(hint_sheet, 2, ["Spieler", "Runde", "Tick", "Kategorie", "Auslösendes Kriterium", "Schwelle", "Beobachtete Fakten", "Review-Aktion"])
    for hint in sorted(hints, key=lambda value: (str(value.get("observed_facts", {}).get("player", "")), int(value.get("observed_facts", {}).get("round_number", 0)), int(value.get("observed_facts", {}).get("first_tick", 0)))):
        facts = hint.get("observed_facts") if isinstance(hint.get("observed_facts"), dict) else {}; threshold = hint.get("threshold") if isinstance(hint.get("threshold"), dict) else {}
        hint_sheet.append([facts.get("player") or "Unbekannt", facts.get("round_number") or "—", facts.get("first_tick") or "—", hint.get("category") or "review_hint", hint.get("criterion_label") or hint.get("criterion_id") or "Unbekannt", f"{threshold.get('field', '—')} {threshold.get('operator', '')} {threshold.get('value', '')}".strip(), hint.get("message") or "Vorhandene Fakten im Kontext prüfen.", "Tactical Replay öffnen; Originaldemo mit Tick im CS2-Kontext prüfen"])
    if not hints: hint_sheet.append(["Keine aktiven Hinweise", "—", "—", "—", "—", "—", "Für dieses Profil und diese Demo wurden keine bewertbaren Hinweise erzeugt.", "—"])
    _table(hint_sheet, "ReviewHintsTable", 2, max(3, len(hints) + 2), 8); _widths(hint_sheet, [24, 13, 13, 26, 28, 22, 48, 48])

    _title(quality_sheet, "Datenqualität · was verfügbar ist und was fehlt", "D"); _header(quality_sheet, 3, ["Status", "Einordnung"])
    user_view = analysis.get("user_view") if isinstance(analysis.get("user_view"), dict) else {}; assessment = user_view.get("assessment") if isinstance(user_view.get("assessment"), dict) else {}
    quality_sheet.append(["Status", (analysis.get("data_quality") or {}).get("status", "unbekannt")]); quality_sheet.append(["Einordnung", assessment.get("status", "Nicht prüfbar / unbekannt")]); quality_sheet.append(["Empfehlung", assessment.get("action", "Keine automatische Interpretation.")])
    _header(quality_sheet, 8, ["Datenkanal", "Bedeutung"])
    available = analysis.get("available_channels") or []
    for channel in available or ["Keine verfügbare Quelle gemeldet."]: quality_sheet.append([channel, "Vom Analyzer geliefert" if available else "Keine Verfügbarkeit gemeldet"])
    quality_sheet.freeze_panes = "A4"; _widths(quality_sheet, [32, 58, 18, 18])

    _title(info_sheet, "Info · Bericht und Interpretation", "F"); _header(info_sheet, 3, ["Feld", "Wert"])
    for value in (("Analyzer-Schema", analysis["schema"]), ("Report-Version", EXCEL_REPORT_SCHEMA), ("Bewertungsprofil", profile.get("label", "Nicht verfügbar")), ("Profil-/Kriterienversion", f"{profile.get('id', '—')}/{profile.get('version', '—')}"), ("Erstellt", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")), ("Map", analysis.get("map_name") or "Unbekannt"), ("Quell-ID", f"{str(analysis.get('source_sha256', ''))[:12]}…"), ("Originaldemo", "Nicht eingebettet; kein lokaler Pfad enthalten")):
        info_sheet.append(value)
    info_sheet.merge_cells("A13:F17"); info_sheet["A13"] = f"{profile.get('disclaimer', 'Automatische Hinweise sind keine Cheat-Nachweise.')} Dieser Bericht fasst vorhandene Analyzer-Ereignisse zusammen; fehlende Datenquellen sind kein negativer Befund."
    info_sheet["A13"].fill = PatternFill("solid", fgColor=_AMBER); info_sheet["A13"].alignment = Alignment(wrap_text=True, vertical="top"); _widths(info_sheet, [28, 58, 14, 14, 14, 14])

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    alignment = copy(cell.alignment)
                    alignment.vertical = "top"; alignment.wrap_text = alignment.wrap_text or sheet.title in {"Multi-Kills & Szenen", "Review-Hinweise"}
                    cell.alignment = alignment
    output_path.parent.mkdir(parents=True, exist_ok=True); workbook.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description="Create a portable Improve Yourself Analyzer Excel report")
    parser.add_argument("analysis", type=Path); parser.add_argument("--output", type=Path, default=Path("results/analyzer-report.xlsx"))
    args = parser.parse_args()
    try: print(build_excel_report(args.analysis, args.output))
    except (OSError, ValueError, json.JSONDecodeError) as error: parser.error(str(error))
    return 0


if __name__ == "__main__": raise SystemExit(main())
